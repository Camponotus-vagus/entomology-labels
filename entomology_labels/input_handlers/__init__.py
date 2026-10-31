"""
Input handlers for various file formats.

Supports: Excel (.xlsx, .xls), CSV, TXT, DOCX, JSON, YAML
"""

import json
import logging
import math
import os
import re
from pathlib import Path
from typing import Any, List, Mapping, Union

from ..config import MAX_COPIES_PER_ENTRY, MAX_FILE_SIZE_BYTES
from ..label_generator import Label, expand_label
from ..sites import has_sites, labels_from_rows, parse_sites

logger = logging.getLogger(__name__)


def _validate_file_path(file_path: Union[str, Path]) -> Path:
    """Check that a path points at a readable file of a workable size.

    This is a usability and resource guard, not a sandbox: the path is
    resolved and accepted wherever it points, which is the correct behaviour
    for a desktop tool whose user picks their own files. It is deliberately
    not confined to a directory and does not reject '..' — callers that need
    confinement must enforce it themselves.

    Args:
        file_path: Path to validate

    Returns:
        Resolved Path object

    Raises:
        FileNotFoundError: If file doesn't exist
        ValueError: If the path is not a file or the file is too large
        PermissionError: If file is not readable
    """
    path = Path(file_path).resolve()

    # Check if file exists
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    # Check if it's actually a file (not directory)
    if not path.is_file():
        raise ValueError(f"Not a file: {file_path}")

    # Check file size to prevent DoS
    try:
        file_size = path.stat().st_size
        if file_size > MAX_FILE_SIZE_BYTES:
            raise ValueError(
                f"File too large: {file_size / (1024*1024):.2f}MB "
                f"(max: {MAX_FILE_SIZE_BYTES / (1024*1024):.0f}MB)"
            )
    except OSError as e:
        raise ValueError(f"Cannot access file: {e}")

    # Check read permissions
    if not os.access(path, os.R_OK):
        raise PermissionError(f"No read permission: {file_path}")

    return path


def _parse_count(value, default: int = 1) -> int:
    """Parse and bound a user-supplied copy count.

    The count comes straight from an input file, so it is validated here —
    before any list of that size is materialised — rather than relying on the
    MAX_LABELS_PER_GENERATOR check, which only runs once the loader has
    already built (and paid for) the full list.

    Args:
        value: Raw count value from the input file
        default: Value to use when the count is missing

    Returns:
        Count as an int between 0 and MAX_COPIES_PER_ENTRY

    Raises:
        ValueError: If the count is not numeric, negative, or above the limit
    """
    if value is None or value == "":
        return default

    if isinstance(value, bool):
        raise ValueError(f"Invalid count value: {value!r} (expected a number)")

    try:
        number = float(str(value).strip())
    except (TypeError, ValueError):
        raise ValueError(f"Invalid count value: {value!r} (expected a number)")

    # A spreadsheet may store 5 as 5.0, which is fine; 2.7 copies is not a
    # count, and truncating it would print a number nobody asked for.
    if not math.isfinite(number) or number != int(number):
        raise ValueError(f"Invalid count value: {value!r} (expected a whole number)")

    count = int(number)
    if count < 0:
        raise ValueError(f"Count cannot be negative: {count}")

    if count > MAX_COPIES_PER_ENTRY:
        raise ValueError(
            f"Count {count} exceeds the maximum of {MAX_COPIES_PER_ENTRY} copies per entry"
        )

    return count


def _sanitize_string(value: str) -> str:
    """Sanitize string input by removing dangerous characters.

    Args:
        value: Input string to sanitize

    Returns:
        Sanitized string with null bytes and control characters removed
    """
    if not value:
        return ""

    # Remove null bytes and control characters except newlines, tabs, carriage returns
    return "".join(c for c in str(value) if c in "\n\r\t" or (ord(c) >= 32 and ord(c) != 127))


def load_data(file_path: Union[str, Path]) -> List[Label]:
    """Load label data from a file.

    Automatically detects the file format based on extension and uses
    the appropriate handler.

    Args:
        file_path: Path to the input file

    Returns:
        List of Label objects

    Raises:
        ValueError: If the file format is not supported
        FileNotFoundError: If the file does not exist
    """
    # Validate and secure the file path
    path = _validate_file_path(file_path)
    logger.info(f"Loading data from: {path}")

    extension = path.suffix.lower()

    handlers = {
        ".xlsx": load_excel,
        ".xls": load_excel,
        ".csv": load_csv,
        ".txt": load_txt,
        ".docx": load_docx,
        ".json": load_json,
        ".yaml": load_yaml,
        ".yml": load_yaml,
    }

    handler = handlers.get(extension)
    if handler is None:
        supported = ", ".join(handlers.keys())
        raise ValueError(
            f"Unsupported file format: {extension}. " f"Supported formats: {supported}"
        )

    try:
        labels = handler(path)
        logger.info(f"Loaded {len(labels)} labels from {path}")
        return labels
    except Exception as e:
        logger.error(f"Error loading data from {path}: {e}", exc_info=True)
        raise


def load_excel(file_path: Path) -> List[Label]:
    """Load labels from an Excel file (.xlsx, .xls).

    Column headings are matched case-insensitively, and several names are
    accepted per field. Where a file carries more than one of them, the first
    listed wins:

    - location_line1: location_line1, location1, location, loc1
    - location_line2: location_line2, location2, loc2
    - code: code, specimen_code, specimen_id, id
    - date: date, collection_date
    - additional_info: additional_info, notes, info (optional)
    - count: count, quantity, copies, n (optional, duplicates the label)

    The Italian headings localita1, localita2, codice, data, data_raccolta,
    note and quantita are also accepted, for spreadsheets written before the
    project moved to English.
    """
    try:
        import pandas as pd
    except ImportError:
        raise ImportError(
            "pandas and openpyxl are required for Excel support. "
            "Install with: pip install pandas openpyxl"
        )

    # Read every cell as text: left to infer types, pandas turns a numeric
    # column with one blank cell into floats (code 5 printed as "5.0"), drops
    # leading zeros, and reads "NA" (Namibia) as a missing value.
    df = pd.read_excel(file_path, dtype=str, keep_default_na=False)
    return _table_to_labels([list(df.columns)] + df.values.tolist())


def _read_delimited(file_path: Path, delimiter: str) -> List[List[str]]:
    """Read a delimited text file into rows of strings.

    The standard library reader, not pandas: CSV is the format people reach
    for first, and it should not need the optional spreadsheet extra. Every
    cell stays text exactly as written, as load_excel also ensures.

    utf-8-sig accepts the byte-order mark spreadsheet programs put at the
    start of an export. Excel's plain "CSV" export on Windows is cp1252,
    which is tried if the file is not valid UTF-8.
    """
    import csv

    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with file_path.open(encoding=encoding, newline="") as handle:
                # strict: an unterminated quote is an error, not a field that
                # silently swallows the rest of the file.
                reader = csv.reader(handle, delimiter=delimiter, strict=True)
                try:
                    return [row for row in reader]
                except csv.Error as e:
                    raise ValueError(f"{file_path.name}, line {reader.line_num}: {e}")
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{file_path.name} is neither UTF-8 nor Windows-1252 text")


def load_csv(file_path: Path) -> List[Label]:
    """Load labels from a CSV file.

    Same column expectations as Excel files.
    Supports both comma and semicolon delimiters.
    """
    # Parse as comma-separated, and only retry with semicolons when the
    # header collapses into a single column (a European-locale export).
    rows = _read_delimited(file_path, ",")
    if rows and len(rows[0]) == 1:
        rows = _read_delimited(file_path, ";")
    return _table_to_labels(rows)


def load_txt(file_path: Path) -> List[Label]:
    """Load labels from a TXT file.

    Supports multiple formats:
    1. Tab-separated values (TSV)
    2. Key-value pairs per label (separated by blank lines)
    3. Simple format: each group of 4-5 lines is a label

    Format 2 example:
        location1: Italia, Trentino Alto Adige,
        location2: Giustino (TN), Vedretta d'Amola
        code: N1
        date: 15.vi.2024

        location1: Italia, Trentino Alto Adige,
        location2: Giustino (TN), Vedretta d'Amola
        code: N2
        date: 15.vi.2024
    """
    content = file_path.read_text(encoding="utf-8-sig")
    lines = content.strip().split("\n")

    # Detect format
    if "\t" in lines[0] and ":" not in lines[0]:
        # TSV format
        rows = _read_delimited(file_path, "\t")
        if _has_known_column(rows[0]):
            return _table_to_labels(rows)

    # Key-value format only if a line actually starts with a known key. A bare
    # colon is not enough: a time ("12:30") in the simple format would
    # otherwise switch parsers and silently lose every field.
    if any(_KEY_VALUE_LINE.match(line) for line in lines):
        return _parse_key_value_txt(content)

    # Simple line-based format
    return _parse_simple_txt(lines)


_KEY_VALUE_KEYS = (
    "location1",
    "location_line1",
    "località1",
    "localita1",
    "location2",
    "location_line2",
    "località2",
    "localita2",
    "code",
    "codice",
    "specimen_code",
    "date",
    "data",
    "collection_date",
    "additional_info",
    "notes",
    "note",
    "count",
    "quantity",
    "quantità",
    "quantita",
)
_KEY_VALUE_LINE = re.compile(
    r"^\s*(?:" + "|".join(re.escape(k) for k in _KEY_VALUE_KEYS) + r")\s*:",
    re.IGNORECASE,
)


def _parse_key_value_txt(content: str) -> List[Label]:
    """Parse key-value formatted text."""
    labels = []
    blocks = re.split(r"\n\s*\n", content.strip())

    for block in blocks:
        if not block.strip():
            continue

        data = {}
        for line in block.strip().split("\n"):
            if ":" in line:
                key, value = line.split(":", 1)
                data[key.strip().lower()] = value.strip()

        if data:
            label = Label(
                location_line1=data.get(
                    "location1",
                    data.get("location_line1", data.get("località1", data.get("localita1", ""))),
                ),
                location_line2=data.get(
                    "location2",
                    data.get("location_line2", data.get("località2", data.get("localita2", ""))),
                ),
                code=data.get("code", data.get("codice", data.get("specimen_code", ""))),
                date=data.get("date", data.get("data", data.get("collection_date", ""))),
                additional_info=data.get(
                    "additional_info", data.get("notes", data.get("note", ""))
                ),
            )

            # Handle count/quantity for duplicates
            count = _parse_count(
                data.get(
                    "count",
                    data.get("quantity", data.get("quantità", data.get("quantita", 1))),
                )
            )
            labels.extend(expand_label(label, count))

    return labels


def _parse_simple_txt(lines: List[str]) -> List[Label]:
    """Parse simple line-based format (4-5 lines per label)."""
    labels = []
    current_lines = []

    for line in lines:
        line = line.strip()
        if not line:
            if current_lines:
                label = _lines_to_label(current_lines)
                if not label.is_empty():
                    labels.append(label)
                current_lines = []
        else:
            current_lines.append(line)

    # Don't forget the last label
    if current_lines:
        label = _lines_to_label(current_lines)
        if not label.is_empty():
            labels.append(label)

    return labels


def _lines_to_label(lines: List[str]) -> Label:
    """Convert a list of lines to a Label."""
    return Label(
        location_line1=lines[0] if len(lines) > 0 else "",
        location_line2=lines[1] if len(lines) > 1 else "",
        code=lines[2] if len(lines) > 2 else "",
        date=lines[3] if len(lines) > 3 else "",
        additional_info=lines[4] if len(lines) > 4 else "",
    )


def load_docx(file_path: Path) -> List[Label]:
    """Load labels from a Word document (.docx).

    Supports two formats:
    1. Table format: Each row is a label with columns for each field
    2. Paragraph format: Labels separated by blank paragraphs
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required for DOCX support. " "Install with: pip install python-docx"
        )

    doc = Document(file_path)
    labels = []

    # Try table format first
    for table in doc.tables:
        if not table.rows:
            continue
        headers = [cell.text.strip().lower() for cell in table.rows[0].cells]

        for row in table.rows[1:]:
            data = {
                headers[i]: cell.text.strip()
                for i, cell in enumerate(row.cells)
                if i < len(headers)
            }
            label = Label.from_dict(data)
            if not label.is_empty():
                labels.append(label)

    # If no tables, try paragraph format
    if not labels:
        paragraphs = [p.text.strip() for p in doc.paragraphs]
        current_lines = []

        for para in paragraphs:
            if not para:
                if current_lines:
                    label = _lines_to_label(current_lines)
                    if not label.is_empty():
                        labels.append(label)
                    current_lines = []
            else:
                current_lines.append(para)

        if current_lines:
            label = _lines_to_label(current_lines)
            if not label.is_empty():
                labels.append(label)

    return labels


def _structured_to_labels(data: Any, items: Any) -> List[Label]:
    """Build labels from a parsed JSON or YAML document.

    Handles both shapes with one code path: the long-standing flat list of
    label rows, and the newer form where a file declares reusable sites and
    defaults. A file with no ``sites:`` or ``defaults:`` key takes exactly the
    old path, so existing files are unaffected.

    Args:
        data: The whole parsed document
        items: The label rows extracted from it

    Returns:
        Labels, with count/quantity already expanded

    Raises:
        ValueError: If a row names a site the file does not declare
    """
    if items is None:
        items = []
    if not isinstance(items, list):
        raise ValueError("labels must be a list of rows")
    for index, row in enumerate(items, start=1):
        if not isinstance(row, dict):
            raise ValueError(f"row {index} is not a mapping of field to value")

    if has_sites(data):
        sites = parse_sites(data.get("sites"))
        defaults = data.get("defaults") or {}
        if not isinstance(defaults, dict):
            raise ValueError("'defaults' must be a mapping of field to value")
        resolved = labels_from_rows(items, sites, defaults)
        counts = [_parse_count(_row_count(row, defaults)) for row in items if isinstance(row, dict)]
    else:
        resolved = [Label.from_dict(item) for item in items]
        counts = [_parse_count(_row_count(row, {})) for row in items]

    labels: List[Label] = []
    for label, count in zip(resolved, counts):
        labels.extend(expand_label(label, count))
    return labels


def _row_count(row: Any, defaults: Mapping[str, Any]) -> Any:
    """Return the copy count for a row, falling back to the file defaults."""
    if not isinstance(row, dict):
        return 1
    for key in ("count", "quantity"):
        if row.get(key) is not None:
            return row[key]
    for key in ("count", "quantity"):
        if defaults.get(key) is not None:
            return defaults[key]
    return 1


def load_json(file_path: Path) -> List[Label]:
    """Load labels from a JSON file.

    Expected format:
    {
        "labels": [
            {
                "location_line1": "Italia, Trentino Alto Adige,",
                "location_line2": "Giustino (TN), Vedretta d'Amola",
                "code": "N1",
                "date": "15.vi.2024",
                "count": 5  // optional, creates duplicates
            },
            ...
        ]
    }

    Or simply an array:
    [
        {"location_line1": "...", ...},
        ...
    ]
    """
    content = file_path.read_text(encoding="utf-8-sig")
    data = json.loads(content)

    if isinstance(data, dict):
        items = data.get("labels", data.get("data", []))
    else:
        items = data

    return _structured_to_labels(data, items)


def load_yaml(file_path: Path) -> List[Label]:
    """Load labels from a YAML file.

    Same structure as JSON format.
    """
    try:
        import yaml
    except ImportError:
        raise ImportError(
            "PyYAML is required for YAML support. " "Install with: pip install pyyaml"
        )

    content = file_path.read_text(encoding="utf-8-sig")
    data = yaml.safe_load(content)

    if isinstance(data, dict):
        items = data.get("labels", data.get("data", []))
    else:
        items = data

    return _structured_to_labels(data, items)


# Accepted column names per field, in priority order. English names are
# listed first so that a file carrying both an English and an Italian
# heading resolves to the English one; the Italian aliases are kept for
# compatibility with spreadsheets written before the project moved to
# English, and removing them would break those files.
_COLUMN_MAP = {
    "location_line1": [
        "location_line1",
        "location1",
        "location",
        "loc1",
        "località1",
        "localita1",
    ],
    "location_line2": [
        "location_line2",
        "location2",
        "loc2",
        "località2",
        "localita2",
    ],
    "code": ["code", "specimen_code", "specimen_id", "id", "codice"],
    "date": ["date", "collection_date", "data_raccolta", "data"],
    "additional_info": ["additional_info", "notes", "info", "note"],
    "count": ["count", "quantity", "copies", "n", "quantità", "quantita"],
    "species": ["species", "taxon", "determination", "specie", "determinazione"],
    "coordinates": ["coordinates", "coords", "latlon", "coordinate"],
    "elevation": ["elevation", "altitude", "quota", "altitudine"],
    "collector": ["collector", "leg", "legit", "raccoglitore"],
    "determiner": ["determiner", "det", "determinatore"],
    "site": ["site", "site_id", "sito"],
}


def _has_known_column(headers: List[str]) -> bool:
    """Whether a header row names at least one field this tool reads."""
    names = {str(h).strip().lower() for h in headers}
    return any(name in names for aliases in _COLUMN_MAP.values() for name in aliases)


def _table_to_labels(rows: List[List[Any]]) -> List[Label]:
    """Convert a table (header row first, then data rows) to Labels.

    Args:
        rows: The header row followed by data rows; cells may be any type
            and missing trailing cells are treated as blank

    Returns:
        Labels, with count/quantity already expanded

    Raises:
        ValueError: If two headings collide or a count is invalid
    """
    if not rows:
        return []

    headers = [str(c).strip().lower() for c in rows[0]]
    duplicated = sorted({h for h in headers if h and headers.count(h) > 1})
    if duplicated:
        # "Code" and "code" collapse to one name, and reading either would be
        # a guess.
        raise ValueError(f"Duplicate column heading(s): {', '.join(duplicated)}")

    columns = {}
    for field, aliases in _COLUMN_MAP.items():
        for name in aliases:
            if name in headers:
                columns[field] = headers.index(name)
                break

    labels = []
    for row in rows[1:]:
        data = {}
        for field, index in columns.items():
            value = row[index] if index < len(row) else None
            # Only missing values become empty. Testing truthiness here
            # would also discard a legitimate 0 (a specimen coded "0").
            if value is None or (hasattr(value, "__float__") and str(value) == "nan"):
                data[field] = ""
            else:
                data[field] = str(value).strip()

        label = Label(**{name: data.get(name, "") for name in Label.CONTENT_FIELDS if name in data})

        if not label.is_empty():
            labels.extend(expand_label(label, _parse_count(data.get("count"))))

    return labels
